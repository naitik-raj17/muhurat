"""Guest import (CSV/XLSX), listing and RSVP updates."""

from __future__ import annotations

import csv
import io
import uuid
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

import phonenumbers
from email_validator import EmailNotValidError, validate_email
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.crypto import pii_hash
from app.db.models import Event, Guest, GuestRsvp
from app.domain.enums import SUPPORTED_LANGUAGES, GuestSide, RsvpStatus
from app.domain.errors import Forbidden, ValidationFailed
from app.domain.guests import GroupingInput, detect_groups
from app.domain.rbac import Principal
from app.domain.schemas import (
    GuestImportReport,
    GuestPage,
    GuestRsvpOut,
    GuestWithRsvps,
    ImportRowError,
    RsvpUpdate,
)
from app.repositories.repos import (
    EventRepository,
    GuestRepository,
    RsvpRepository,
    WeddingRepository,
)

MAX_IMPORT_ROWS = 5000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
DEFAULT_REGION = "IN"

# Accepted header spellings -> canonical field.
_HEADER_ALIASES = {
    "name": "name",
    "full name": "name",
    "guest name": "name",
    "phone": "phone",
    "mobile": "phone",
    "phone number": "phone",
    "whatsapp": "phone",
    "email": "email",
    "e-mail": "email",
    "language": "language",
    "lang": "language",
    "side": "side",
    "family": "family",
    "group": "family",
    "household": "family",
    "address": "address",
    "city": "city",
    "relation": "relation",
    "distance_km": "distance_km",
    "distance km": "distance_km",
    "vip": "vip",
}
_SIDE_ALIASES = {
    "bride": "bride",
    "ladkiwale": "bride",
    "groom": "groom",
    "ladkewale": "groom",
    "both": "both",
    "common": "both",
}
_LANG_ALIASES = {"english": "en", "hindi": "hi", "हिंदी": "hi", "en": "en", "hi": "hi"}


# --------------------------------------------------------------------------- parsing


def normalise_phone(raw: str, region: str = DEFAULT_REGION) -> str:
    """E.164 or ValueError. Indian numbers without a country code are assumed (+91)."""
    try:
        num = phonenumbers.parse(raw, region)
    except phonenumbers.NumberParseException as exc:
        raise ValueError("not a phone number") from exc
    if not phonenumbers.is_valid_number(num):
        raise ValueError("invalid phone number")
    return phonenumbers.format_number(num, phonenumbers.PhoneNumberFormat.E164)


def _rows_from_csv(data: bytes) -> Iterator[dict[str, str]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    for row in reader:
        yield {(k or "").strip(): (v or "").strip() for k, v in row.items()}


def _rows_from_xlsx(data: bytes) -> Iterator[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return
    keys = [str(h or "").strip() for h in header]
    for values in rows:
        if values is None or all(v is None for v in values):
            continue
        yield {k: ("" if v is None else str(v).strip()) for k, v in zip(keys, values, strict=False)}


def parse_file(filename: str, data: bytes) -> list[dict[str, str]]:
    if len(data) > MAX_IMPORT_BYTES:
        raise ValidationFailed("File is larger than 5 MB")
    name = filename.lower()
    if name.endswith(".csv"):
        raw_rows = list(_rows_from_csv(data))
    elif name.endswith(".xlsx"):
        try:
            raw_rows = list(_rows_from_xlsx(data))
        except Exception as exc:  # openpyxl raises many types for corrupt files
            raise ValidationFailed("Could not read the Excel file") from exc
    else:
        raise ValidationFailed("Upload a .csv or .xlsx file")
    if len(raw_rows) > MAX_IMPORT_ROWS:
        raise ValidationFailed(f"At most {MAX_IMPORT_ROWS} rows per import")
    canonical: list[dict[str, str]] = []
    for row in raw_rows:
        out: dict[str, str] = {}
        for key, value in row.items():
            field = _HEADER_ALIASES.get(key.strip().lower())
            if field and value:
                out[field] = value
        canonical.append(out)
    if raw_rows and not any("name" in r for r in canonical):
        raise ValidationFailed("No 'name' column found. Expected headers: name, phone, email, side")
    return canonical


@dataclass
class CleanRow:
    row: int  # 1-based data row, as shown in a spreadsheet (header is row 1, so data starts at 2)
    name: str
    phone: str | None
    email: str | None
    language: str
    side: str
    family: str | None
    address: str | None
    city: str | None
    relation: str | None
    distance_km: float | None
    is_vip: bool


def validate_row(n: int, raw: dict[str, str]) -> tuple[CleanRow | None, list[ImportRowError]]:
    errors: list[ImportRowError] = []

    def err(field: str, msg: str) -> None:
        errors.append(ImportRowError(row=n, field=field, message=msg, value=raw.get(field)))

    name = raw.get("name", "").strip()
    if not name:
        err("name", "name is required")
    elif len(name) > 200:
        err("name", "name is longer than 200 characters")

    phone = None
    if raw.get("phone"):
        try:
            phone = normalise_phone(raw["phone"])
        except ValueError as exc:
            err("phone", str(exc))
    email = None
    if raw.get("email"):
        try:
            email = validate_email(raw["email"], check_deliverability=False).normalized.lower()
        except EmailNotValidError:
            err("email", "invalid email address")
    if not raw.get("phone") and not raw.get("email"):
        err("phone", "a phone number or email is required to invite this guest")

    side = _SIDE_ALIASES.get(raw.get("side", "").strip().lower(), "")
    if not side:
        err("side", "side must be bride, groom or both")

    language = _LANG_ALIASES.get(raw.get("language", "en").strip().lower(), "")
    if language not in SUPPORTED_LANGUAGES:
        err("language", f"language must be one of {', '.join(SUPPORTED_LANGUAGES)}")

    distance = None
    if raw.get("distance_km"):
        try:
            distance = float(raw["distance_km"])
            if distance < 0:
                raise ValueError
        except ValueError:
            err("distance_km", "distance_km must be a non-negative number")

    if errors:
        return None, errors
    return (
        CleanRow(
            row=n,
            name=name,
            phone=phone,
            email=email,
            language=language,
            side=side,
            family=raw.get("family"),
            address=raw.get("address"),
            city=raw.get("city"),
            relation=raw.get("relation"),
            distance_km=distance,
            is_vip=raw.get("vip", "").strip().lower() in {"1", "yes", "y", "true", "vip"},
        ),
        [],
    )


# --------------------------------------------------------------------------- import


async def import_guests(
    session: AsyncSession, wedding_id: uuid.UUID, filename: str, data: bytes
) -> GuestImportReport:
    await WeddingRepository(session).get(wedding_id)
    raw_rows = parse_file(filename, data)
    guests = GuestRepository(session)

    existing_phone = set()
    existing_email = set()
    stmt = select(Guest.phone_hash, Guest.email_hash).where(
        guests.scope_filter(), Guest.wedding_id == wedding_id
    )
    for ph, em in (await session.execute(stmt)).all():
        if ph:
            existing_phone.add(ph)
        if em:
            existing_email.add(em)

    errors: list[ImportRowError] = []
    clean: list[tuple[CleanRow, str | None, str | None]] = []
    seen_phone: set[str] = set()
    seen_email: set[str] = set()
    dup_file = dup_existing = 0
    for i, raw in enumerate(raw_rows, start=2):
        row, row_errors = validate_row(i, raw)
        if row is None:
            errors.extend(row_errors)
            continue
        ph = pii_hash(row.phone) if row.phone else None
        em = pii_hash(row.email) if row.email else None
        if (ph and ph in existing_phone) or (em and em in existing_email):
            dup_existing += 1
            errors.append(
                ImportRowError(
                    row=i, field="phone" if ph else "email", message="guest already exists"
                )
            )
            continue
        # Same phone or email twice in the file: keep the first, report the rest. Family members
        # sharing one number are common in India, so a shared phone with a different name is kept.
        key_phone = f"{ph}|{row.name.lower()}" if ph else None
        if (key_phone and key_phone in seen_phone) or (em and em in seen_email):
            dup_file += 1
            errors.append(
                ImportRowError(row=i, field="name", message="duplicate of an earlier row in file")
            )
            continue
        if key_phone:
            seen_phone.add(key_phone)
        if em:
            seen_email.add(em)
        clean.append((row, ph, em))

    groups = detect_groups(
        [
            GroupingInput(
                index=pos,
                name=row.name,
                side=row.side,
                phone_hash=ph,
                email_hash=em,
                family_label=row.family,
                address=row.address,
            )
            for pos, (row, ph, em) in enumerate(clean)
        ]
    )
    group_ids: dict[int, uuid.UUID] = {}

    events = await EventRepository(session).for_wedding(wedding_id)
    rsvps = RsvpRepository(session)
    for pos, (row, ph, em) in enumerate(clean):
        group = groups.get(pos)
        guest = guests.add(
            Guest(
                wedding_id=wedding_id,
                name=row.name,
                phone=row.phone,
                phone_hash=ph,
                email=row.email,
                email_hash=em,
                language=row.language,
                side=row.side,
                group_id=group_ids.setdefault(group[0], uuid.uuid4()) if group else None,
                group_label=group[1] if group else None,
                city=row.city,
                distance_km=row.distance_km,
                relation=row.relation,
                is_vip=row.is_vip,
            )
        )
        await session.flush()
        for ev in events:
            rsvps.add(GuestRsvp(guest_id=guest.id, event_id=ev.id, status="pending"))
    await session.flush()

    return GuestImportReport(
        total_rows=len(raw_rows),
        imported=len(clean),
        skipped=len(raw_rows) - len(clean),
        duplicates_in_file=dup_file,
        duplicates_existing=dup_existing,
        groups_detected=len(group_ids),
        errors=errors,
    )


# --------------------------------------------------------------------------- read / RSVP


async def list_guests(
    session: AsyncSession,
    wedding_id: uuid.UUID,
    *,
    search: str | None = None,
    side: GuestSide | None = None,
    rsvp_status: RsvpStatus | None = None,
    limit: int = 50,
    offset: int = 0,
) -> GuestPage:
    await WeddingRepository(session).get(wedding_id)
    guests = GuestRepository(session)
    where: list[Any] = [Guest.wedding_id == wedding_id]
    if search:
        like = f"%{search.strip()}%"
        where.append(or_(Guest.name.ilike(like), Guest.group_label.ilike(like)))
    if side:
        where.append(Guest.side == side.value)
    if rsvp_status:
        where.append(
            Guest.id.in_(select(GuestRsvp.guest_id).where(GuestRsvp.status == rsvp_status.value))
        )
    total = await guests.count(*where)
    page = await guests.list(
        *where, order_by=[Guest.group_label.nulls_last(), Guest.name], limit=limit, offset=offset
    )
    rsvp_rows = await RsvpRepository(session).list(GuestRsvp.guest_id.in_([g.id for g in page]))
    by_guest: dict[uuid.UUID, list[GuestRsvpOut]] = {}
    for r in rsvp_rows:
        by_guest.setdefault(r.guest_id, []).append(GuestRsvpOut.model_validate(r))
    items = []
    for g in page:
        item = GuestWithRsvps.model_validate(g)
        item.rsvps = by_guest.get(g.id, [])
        items.append(item)
    return GuestPage(total=total, items=items)


async def update_rsvp(
    session: AsyncSession,
    principal: Principal,
    wedding_id: uuid.UUID,
    guest_id: uuid.UUID,
    data: RsvpUpdate,
) -> GuestRsvpOut:
    guest = await GuestRepository(session).get(guest_id)
    if guest.wedding_id != wedding_id:
        raise Forbidden("Guest belongs to another wedding")
    event = await EventRepository(session).get(data.event_id)
    if event.wedding_id != wedding_id:
        raise ValidationFailed("Event belongs to another wedding")
    rsvps = RsvpRepository(session)
    existing = await rsvps.list(GuestRsvp.guest_id == guest_id, GuestRsvp.event_id == event.id)
    rsvp = existing[0] if existing else rsvps.add(GuestRsvp(guest_id=guest_id, event_id=event.id))
    rsvp.status = data.status.value
    rsvp.plus_ones = data.plus_ones
    rsvp.responded_at = datetime.now(UTC)
    rsvp.updated_by_actor = principal.actor_type.value
    await session.flush()
    return GuestRsvpOut.model_validate(rsvp)


async def rsvp_counts_by_event(
    session: AsyncSession, wedding_id: uuid.UUID
) -> dict[str, dict[str, int]]:
    repo = RsvpRepository(session)
    stmt = (
        select(Event.id, GuestRsvp.status, func.count())
        .join(Event, Event.id == GuestRsvp.event_id)
        .where(repo.scope_filter(), Event.wedding_id == wedding_id)
        .group_by(Event.id, GuestRsvp.status)
    )
    out: dict[str, dict[str, int]] = {}
    for event_id, status, n in (await session.execute(stmt)).all():
        out.setdefault(str(event_id), {})[status] = int(n)
    return out
